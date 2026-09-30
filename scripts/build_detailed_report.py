from __future__ import annotations

# ruff: noqa: E501
import argparse
import math
from collections.abc import Iterable, Sequence
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

PROJECT_NAME = "NVIDIA Industrial Catalog Extraction Pipeline"
REPORT_DATE = "7 August 2026"

INK = "#172033"
NAVY = "#0B1F33"
BLUE = "#285A8E"
GREEN = "#76B900"
GREEN_DARK = "#4C7F00"
PALE_GREEN = "#EEF7E4"
PALE_BLUE = "#EDF4FA"
PALE_GRAY = "#F3F5F7"
MID_GRAY = "#68717D"
LINE = "#CDD4DC"
WHITE = "#FFFFFF"
AMBER = "#F6B73C"
PALE_AMBER = "#FFF5DC"
RED = "#B33A3A"
PALE_RED = "#FCEAEA"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value.lstrip("#"))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_indent(table, twips: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(twips))
    tbl_ind.set(qn("w:type"), "dxa")


def set_table_width(table, twips: int = 9360) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(twips))
    tbl_w.set(qn("w:type"), "dxa")


def set_grid_widths(table, widths: Sequence[int]) -> None:
    grid_columns = list(table._tbl.tblGrid.gridCol_lst)
    for index, width in enumerate(widths):
        if index >= len(grid_columns):
            grid_column = OxmlElement("w:gridCol")
            table._tbl.tblGrid.append(grid_column)
            grid_columns.append(grid_column)
        grid_columns[index].set(qn("w:w"), str(width))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def disallow_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(twips))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    text_node = OxmlElement("w:t")
    text_node.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend((fld_char_begin, instr_text, fld_char_separate, text_node, fld_char_end))


def set_page_number_start(section, start: int) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def set_run_font(run, name: str = "Calibri", size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def keep(paragraph, *, next_: bool = False, together: bool = False) -> None:
    paragraph.paragraph_format.keep_with_next = next_
    paragraph.paragraph_format.keep_together = together


def add_hyperlink(paragraph, text: str, url: str, color: str = BLUE) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color.lstrip("#"))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((color_el, underline))
    run.append(run_properties)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_alt_text(inline_shape, title: str, description: str) -> None:
    try:
        doc_pr = inline_shape._inline.docPr
        doc_pr.set("title", title)
        doc_pr.set("descr", description)
    except Exception:
        pass


def add_label(document: Document, text: str, color: str = GREEN_DARK) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    keep(p, next_=True)
    r = p.add_run(text.upper())
    set_run_font(r, size=8.5, color=color, bold=True)
    r.font.letter_spacing = Pt(0.4) if hasattr(r.font, "letter_spacing") else None


def add_title(document: Document, text: str, level: int = 1) -> None:
    p = document.add_heading(text, level=level)
    keep(p, next_=True)


def add_body(document: Document, text: str, *, bold_lead: str | None = None) -> None:
    p = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)


def add_bullet(document: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = document.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    keep(p, together=True)
    r = p.add_run(text)
    set_run_font(r)


def add_number(document: Document, text: str) -> None:
    p = document.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    keep(p, together=True)
    r = p.add_run(text)
    set_run_font(r)


def add_callout(document: Document, title: str, text: str, *, kind: str = "info") -> None:
    palette = {
        "info": (PALE_BLUE, BLUE),
        "success": (PALE_GREEN, GREEN_DARK),
        "warning": (PALE_AMBER, "#8A5A00"),
        "danger": (PALE_RED, RED),
    }
    fill, accent = palette[kind]
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_indent(table, 120)
    set_grid_widths(table, (180, 9180))
    set_cell_width(table.cell(0, 0), 180)
    set_cell_width(table.cell(0, 1), 9180)
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_margins(table.cell(0, 0), 80, 120, 80, 120)
    set_cell_margins(table.cell(0, 1), 120, 120, 120, 120)
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + "\n")
    set_run_font(r, size=10.5, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=9.5, color=INK)
    disallow_row_split(table.rows[0])
    set_repeat_table_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(document: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int] | None = None,
              *, font_size: float = 8.8) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_indent(table, 120)
    if widths is None:
        widths = [9360 // len(headers)] * len(headers)
    set_grid_widths(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, color=WHITE, bold=True)
    set_repeat_table_header(table.rows[0])
    disallow_row_split(table.rows[0])
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        fill = WHITE if row_index % 2 == 0 else PALE_GRAY
        for index, value in enumerate(row_values):
            cell = cells[index]
            set_cell_width(cell, widths[index])
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
        disallow_row_split(table.rows[-1])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(document: Document, lines: Iterable[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_indent(table, 120)
    set_grid_widths(table, (9360,))
    cell = table.cell(0, 0)
    set_cell_shading(cell, "17212B")
    set_cell_margins(cell, 120, 120, 120, 120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(lines):
        if idx:
            p.add_run("\n")
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=8.0, color="E8EDF2")
    disallow_row_split(table.rows[0])
    set_repeat_table_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(document: Document, image_path: Path, caption: str, alt_text: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    keep(p, next_=True, together=True)
    shape = p.add_run().add_picture(str(image_path), width=Inches(6.45))
    set_alt_text(shape, caption, alt_text)
    cp = document.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(8)
    keep(cp, together=True)
    r = cp.add_run(caption)
    set_run_font(r, size=8.5, color=MID_GRAY, italic=True)


def find_font(bold: bool = False) -> str:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    return "arial.ttf"


FONT_REGULAR = find_font(False)
FONT_BOLD = find_font(True)


def font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REGULAR, size)


def wrapped_lines(draw: ImageDraw.ImageDraw, text: str, font_obj, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font_obj)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_centered_text(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], text: str, font_obj,
                       fill: str, max_width: int | None = None, line_gap: int = 8) -> None:
    x1, y1, x2, y2 = box
    width = max_width or (x2 - x1 - 40)
    lines = wrapped_lines(draw, text, font_obj, width)
    heights = [draw.textbbox((0, 0), line, font=font_obj)[3] for line in lines]
    total = sum(heights) + line_gap * max(0, len(lines) - 1)
    cursor_y = y1 + (y2 - y1 - total) / 2
    for line, line_height in zip(lines, heights, strict=True):
        bbox = draw.textbbox((0, 0), line, font=font_obj)
        cursor_x = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((cursor_x, cursor_y), line, font=font_obj, fill=fill)
        cursor_y += line_height + line_gap


def draw_box(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str, subtitle: str,
             *, fill: str = WHITE, outline: str = LINE, accent: str = GREEN) -> None:
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=3)
    draw.rounded_rectangle((x1, y1, x1 + 16, y2), radius=8, fill=accent)
    draw_centered_text(draw, (x1 + 35, y1 + 12, x2 - 15, y1 + 70), title, font(25, True), NAVY)
    draw_centered_text(draw, (x1 + 35, y1 + 70, x2 - 15, y2 - 12), subtitle, font(18), MID_GRAY, line_gap=5)


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = BLUE, width: int = 7) -> None:
    draw.line((start, end), fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 18
    left = (end[0] - size * math.cos(angle - math.pi / 6), end[1] - size * math.sin(angle - math.pi / 6))
    right = (end[0] - size * math.cos(angle + math.pi / 6), end[1] - size * math.sin(angle + math.pi / 6))
    draw.polygon([end, left, right], fill=color)


def diagram_canvas(title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (2400, 1350), WHITE)
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 2400, 118), fill=NAVY)
    draw.text((80, 26), title, font=font(42, True), fill=WHITE)
    draw.text((80, 130), subtitle, font=font(22), fill=MID_GRAY)
    draw.rectangle((80, 184, 360, 193), fill=GREEN)
    return image, draw


def create_architecture_diagram(path: Path) -> None:
    image, draw = diagram_canvas(
        "End-to-end evidence pipeline",
        "NVIDIA inference stages are bounded by deterministic routing, validation, lineage, and audit controls.",
    )
    boxes = [
        ((70, 280, 395, 500), "1  PDF intake", "SHA-256 identity\nPyMuPDF page stream\nsource-object inventory"),
        ((475, 280, 800, 500), "2  Page router", "native quality\nlayout complexity\nblank-page signals"),
        ((880, 230, 1225, 450), "3A  Nemotron Parse", "page image → text, tables,\nclasses and normalized bboxes"),
        ((880, 520, 1225, 720), "3B  Native text", "quality-approved blocks\nwith page coordinates"),
        ((1305, 280, 1640, 500), "4  Evidence layer", "deterministic element IDs\nexact excerpts\nfull provider audit"),
        ((1720, 230, 2055, 450), "5  Nemotron Nano", "guided JSON\nmanufacturer + cardinality hints\nproduct candidates"),
        ((1720, 520, 2055, 720), "6  Source guard", "evidence hydration\nPydantic validation\nfail-closed admission"),
        ((2080, 375, 2330, 600), "7  Knowledge base", "SQLite catalog\ntwo-tier RAG\nNemotron Embed"),
    ]
    for box, title, subtitle in boxes:
        accent = GREEN if title[0] in "13457" else BLUE
        fill = PALE_GREEN if title.startswith(("3A", "5", "7")) else WHITE
        draw_box(draw, box, title, subtitle, fill=fill, accent=accent)
    draw_arrow(draw, (395, 390), (475, 390))
    draw_arrow(draw, (800, 370), (880, 340))
    draw_arrow(draw, (800, 420), (880, 620))
    draw_arrow(draw, (1225, 340), (1305, 370))
    draw_arrow(draw, (1225, 620), (1305, 430))
    draw_arrow(draw, (1640, 370), (1720, 340))
    draw_arrow(draw, (1885, 450), (1885, 520))
    draw_arrow(draw, (2055, 620), (2080, 520))
    draw_arrow(draw, (1470, 500), (1470, 970), GREEN_DARK)
    draw_arrow(draw, (1470, 970), (2195, 970), GREEN_DARK)
    draw_arrow(draw, (2195, 970), (2195, 600), GREEN_DARK)
    draw.rounded_rectangle((190, 870, 1210, 1130), radius=24, fill=PALE_BLUE, outline=BLUE, width=3)
    draw.text((235, 905), "RAW EVIDENCE RAG", font=font(28, True), fill=BLUE)
    draw.text((235, 960), "Every nonblank extracted element is indexed before\nproduct admission, preserving searchable evidence even\nwhen a product candidate is rejected.", font=font(23), fill=INK, spacing=10)
    draw.rounded_rectangle((1300, 1040, 2280, 1260), radius=24, fill=PALE_GREEN, outline=GREEN_DARK, width=3)
    draw.text((1345, 1070), "CANONICAL PRODUCT RAG", font=font(28, True), fill=GREEN_DARK)
    draw.text((1345, 1125), "Only source-guarded, validated product identity,\nspecification, and operating-condition records enter\nthe canonical catalog tier.", font=font(23), fill=INK, spacing=10)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(220, 220), optimize=True)


def create_resilience_diagram(path: Path) -> None:
    image, draw = diagram_canvas(
        "Page extraction decision and recovery",
        "Transport success is not semantic success; blank, fallback, and failure outcomes are independently audited.",
    )
    draw_box(draw, (850, 225, 1550, 380), "Inventory page", "text • blocks • images • drawings • annotations • links", fill=PALE_BLUE, accent=BLUE)
    draw_arrow(draw, (1200, 380), (1200, 460))
    draw_box(draw, (850, 460, 1550, 630), "Certified blank?", "Requires an exact zero inventory across every source-object signal", fill=WHITE, accent=GREEN)
    draw_arrow(draw, (850, 545), (580, 545), GREEN_DARK)
    draw_box(draw, (80, 455, 580, 640), "YES — deterministic empty", "No evidence elements\nSkip structured LLM\nWrite stable page audit", fill=PALE_GREEN, accent=GREEN_DARK)
    draw.text((650, 505), "YES", font=font(22, True), fill=GREEN_DARK)
    draw.arrow = None
    draw_arrow(draw, (1200, 630), (1200, 705), BLUE)
    draw.text((1230, 650), "NO", font=font(22, True), fill=BLUE)
    draw_box(draw, (820, 705, 1580, 890), "Route and extract", "Native blocks for clean text; Nemotron Parse for scans, tables, columns, or image-heavy pages", fill=PALE_BLUE, accent=BLUE)
    draw_arrow(draw, (1580, 795), (1790, 795), BLUE)
    draw_box(draw, (1790, 705, 2310, 890), "Parse result usable?", "Supported, nonblank assistant content is required even after HTTP 200", fill=WHITE, accent=AMBER)
    draw_arrow(draw, (2050, 890), (2050, 985), GREEN_DARK)
    draw.text((2080, 915), "YES", font=font(22, True), fill=GREEN_DARK)
    draw_box(draw, (1740, 985, 2320, 1195), "Normalize evidence", "Tagged text/class/bbox or JSON elements\nPreserve raw provider payload\nCreate deterministic IDs", fill=PALE_GREEN, accent=GREEN_DARK)
    draw_arrow(draw, (1790, 760), (1610, 540), AMBER)
    draw_box(draw, (1640, 390, 2310, 590), "NO — retry up to 8 attempts", "Retry transport errors and successful-but-empty/unsupported assistant payloads; retain every response", fill=PALE_AMBER, accent="#8A5A00")
    draw_arrow(draw, (1840, 390), (1840, 320), RED)
    draw_box(draw, (1510, 170, 2050, 320), "Exhausted: native evidence passes?", "Only exact HTTP-200 unsupported-content exhaustion is eligible", fill=PALE_RED, accent=RED)
    draw_arrow(draw, (1510, 245), (1480, 245), GREEN_DARK)
    draw.text((1505, 195), "YES", font=font(22, True), fill=GREEN_DARK)
    draw_box(draw, (930, 170, 1480, 300), "Audited native fallback", "Keep original VLM route\nUse quality-approved native blocks", fill=PALE_GREEN, accent=GREEN_DARK)
    draw_arrow(draw, (2050, 245), (2090, 245), RED)
    draw.text((2058, 198), "NO", font=font(20, True), fill=RED)
    draw.rounded_rectangle((2090, 150, 2380, 340), radius=18, fill=PALE_RED, outline=RED, width=3)
    draw_centered_text(draw, (2105, 164, 2365, 326), "FAIL CLOSED — unresolved nonblank page; no catalog admission", font(18, True), RED, max_width=235, line_gap=4)
    draw.rounded_rectangle((110, 1030, 1510, 1245), radius=24, fill=PALE_GRAY, outline=LINE, width=3)
    draw.text((155, 1065), "Every terminal page outcome records", font=font(27, True), fill=NAVY)
    draw.text((155, 1120), "routing signals • reason codes • extraction method • provider payload • attempts • model identity • page-level audit", font=font(21), fill=INK)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(220, 220), optimize=True)


def create_topology_diagram(path: Path) -> None:
    image, draw = diagram_canvas(
        "Verified NVIDIA server deployment",
        "Three isolated OpenAI-compatible inference endpoints support a sequential, bounded runner.",
    )
    draw.rounded_rectangle((80, 250, 540, 1160), radius=25, fill=PALE_GRAY, outline=LINE, width=3)
    draw.text((120, 285), "ORCHESTRATION", font=font(28, True), fill=NAVY)
    draw_box(draw, (120, 365, 500, 555), "Pipeline runner", "Python 3.12\nbounded document/page limits\natomic checkpoints", fill=WHITE, accent=BLUE)
    draw_box(draw, (120, 650, 500, 840), "Monitoring", "endpoint health\nGPU ownership\nrun summaries + SQLite checks", fill=WHITE, accent=GREEN)
    draw_box(draw, (120, 930, 500, 1105), "Storage", "/data/industrial-data-corps\nbenchmarks • manifests • releases", fill=WHITE, accent=GREEN_DARK)
    draw.rounded_rectangle((650, 250, 2300, 1160), radius=25, fill=WHITE, outline=LINE, width=3)
    draw.text((700, 285), "8 × NVIDIA A100-SXM4 80 GB  |  640 GB aggregate VRAM", font=font(30, True), fill=NAVY)
    cards = [
        ((720, 385, 1160, 680), "GPU 1", "Nemotron Parse v1.2", "Port 8001\nDocument VLM\nvLLM max sequences: 8", GREEN),
        ((1260, 385, 1700, 680), "GPU 3", "Nemotron Nano 8B", "Port 8003\nStructured LLM\nvLLM max sequences: 16", BLUE),
        ((1800, 385, 2240, 680), "GPU 5", "Nemotron Embed 8B", "Port 8004\n4,096-D vectors\nvLLM max sequences: 16", GREEN_DARK),
    ]
    for box, gpu, model, details, accent in cards:
        x1, y1, x2, y2 = box
        draw.rounded_rectangle(box, radius=22, fill=PALE_GREEN if accent != BLUE else PALE_BLUE, outline=accent, width=4)
        draw.text((x1 + 30, y1 + 25), gpu, font=font(31, True), fill=accent)
        draw.text((x1 + 30, y1 + 85), model, font=font(24, True), fill=NAVY)
        draw.multiline_text((x1 + 30, y1 + 145), details, font=font(22), fill=INK, spacing=12)
    draw.rounded_rectangle((720, 790, 2240, 1075), radius=22, fill=PALE_GRAY, outline=LINE, width=3)
    draw.text((760, 825), "GPU allocation guardrails", font=font(28, True), fill=NAVY)
    draw.text((760, 885), "Reserved: GPU 0 and GPU 2 — never adopted, reset, or stopped by this project", font=font(22), fill=RED)
    draw.text((760, 935), "Available for measured expansion: GPU 4, GPU 6, GPU 7", font=font(22), fill=GREEN_DARK)
    draw.text((760, 985), "Launcher adopts a service only after PID, model, port, GPU ownership, and health match", font=font(22), fill=INK)
    draw_arrow(draw, (500, 460), (720, 520), BLUE)
    draw_arrow(draw, (500, 735), (720, 910), GREEN_DARK)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, dpi=(220, 220), optimize=True)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.24 if style_name != "List Bullet 2" else 0.46)
        style.paragraph_format.first_line_indent = Inches(-0.16)
        style.paragraph_format.space_after = Pt(3)

    heading_settings = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, NAVY, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_settings.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Small Caps Label" not in styles:
        label_style = styles.add_style("Small Caps Label", WD_STYLE_TYPE.PARAGRAPH)
        label_style.font.name = "Calibri"
        label_style.font.size = Pt(8.5)
        label_style.font.bold = True
        label_style.font.color.rgb = rgb(GREEN_DARK)


def configure_section(section, *, cover: bool = False) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78 if cover else 0.8)
    section.bottom_margin = Inches(0.68 if cover else 0.72)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def configure_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("NVIDIA INDUSTRIAL CATALOG EXTRACTION  •  TECHNICAL REPORT")
    set_run_font(r, size=7.5, color=MID_GRAY, bold=True)
    fp = section.footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    table = fp._parent.add_table(rows=1, cols=2, width=Inches(6.5))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_grid_widths(table, (8064, 1296))
    set_cell_width(table.cell(0, 0), 8064)
    set_cell_width(table.cell(0, 1), 1296)
    for cell in table.rows[0].cells:
        set_cell_margins(cell, 0, 0, 0, 0)
    set_repeat_table_header(table.rows[0])
    p_left = table.cell(0, 0).paragraphs[0]
    r_left = p_left.add_run(f"Architecture and implementation report  •  {REPORT_DATE}")
    set_run_font(r_left, size=7.5, color=MID_GRAY)
    p_right = table.cell(0, 1).paragraphs[0]
    add_page_number(p_right)
    for run in p_right.runs:
        set_run_font(run, size=7.5, color=MID_GRAY)


def cover_metric_table(document: Document) -> None:
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_indent(table, 200)
    set_grid_widths(table, (4680, 4680))
    metrics = [
        ("8 × A100 80 GB", "Verified NVIDIA GPU server"),
        ("220 PDFs", "Transferred and inventory-verified"),
        ("17,837 pages", "Two pinned corpus profiles"),
        ("140+ tests passed", "Local and remote verification"),
    ]
    for index, (value, label) in enumerate(metrics):
        row = index // 2
        col = index % 2
        cell = table.cell(row, col)
        set_cell_width(cell, 4680)
        set_cell_shading(cell, PALE_GRAY if index % 3 else PALE_GREEN)
        set_cell_margins(cell, 150, 200, 150, 200)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(value + "\n")
        set_run_font(r, size=17, color=GREEN_DARK if index % 3 == 0 else NAVY, bold=True)
        r2 = p.add_run(label)
        set_run_font(r2, size=8.5, color=MID_GRAY)
        disallow_row_split(table.rows[row])
    set_repeat_table_header(table.rows[0])


def add_cover(document: Document) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(70)
    r = p.add_run("ENGINEERING REPORT  /  SOFTWARE RELEASE 0.1.0")
    set_run_font(r, size=9, color=GREEN_DARK, bold=True)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("NVIDIA Industrial\nCatalog Extraction Pipeline")
    set_run_font(r, size=31, color=NAVY, bold=True)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("How it works, system architecture, evidence controls, validation results, and the path to full-corpus production")
    set_run_font(r, size=14, color=BLUE)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Prepared for the Industrial Data Catalogs program\n")
    set_run_font(r, size=10.5, color=INK, bold=True)
    r2 = p.add_run(f"Status as of {REPORT_DATE}  •  Local + NVIDIA server implementation")
    set_run_font(r2, size=9.5, color=MID_GRAY)

    cover_metric_table(document)
    document.add_paragraph().paragraph_format.space_after = Pt(14)

    add_callout(
        document,
        "Status boundary",
        "The architecture and bounded production validation are complete. The full 17,837-page corpus has not yet been processed; manifest-pinned page-window sharding and deterministic merge controls are required before that claim is valid.",
        kind="warning",
    )
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("Public repository")
    set_run_font(r, size=8, color=MID_GRAY, bold=True)
    r2 = p.add_run("\nRepository-relative code, tests, documentation, and report artifacts")
    set_run_font(r2, name="Consolas", size=7.7, color=INK)


def add_report_map(document: Document) -> None:
    add_label(document, "Report map")
    add_title(document, "Contents and reading guide", 1)
    rows = [
        ("1–3", "Decision context", "Executive result, design goals, infrastructure, and corpus scope"),
        ("4–7", "Architecture", "Routing, VLM extraction, structured parsing, evidence guards, and resilience"),
        ("8–10", "Data and operations", "Canonical schema, two-tier RAG, checkpoints, deployment, and monitoring"),
        ("11–12", "Proof", "Validation runs, incident replay, integrity results, and throughput interpretation"),
        ("13–14", "Production path", "Risks, manifest sharding, deterministic merge, and delivery roadmap"),
        ("Appendices", "Reference", "Implementation map, output artifacts, glossary, and authoritative model sources"),
    ]
    add_table(document, ("Sections", "Theme", "What the reader gets"), rows, (1350, 2400, 5610), font_size=9.0)
    add_callout(document, "Fast path for decision-makers", "Read the Executive Summary, Verified Results, Risk Register, and Production Scale Roadmap. The remaining sections provide implementation-level traceability.", kind="info")
    add_callout(
        document,
        "Independent research and licensing boundary",
        "This is independent research software and is not an NVIDIA product or endorsement. Repository code is offered under Apache-2.0; that license does not grant rights to NVIDIA model weights, third-party dependencies, or source catalog content. Review each upstream model, dependency, and data license before use or redistribution.",
        kind="warning",
    )


def build_report(output_path: Path, asset_dir: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    asset_dir.mkdir(parents=True, exist_ok=True)
    architecture_path = asset_dir / "architecture_overview.png"
    resilience_path = asset_dir / "resilience_flow.png"
    topology_path = asset_dir / "deployment_topology.png"
    create_architecture_diagram(architecture_path)
    create_resilience_diagram(resilience_path)
    create_topology_diagram(topology_path)

    document = Document()
    configure_styles(document)
    configure_section(document.sections[0], cover=True)
    document.sections[0].different_first_page_header_footer = True
    document.core_properties.title = PROJECT_NAME
    document.core_properties.subject = "Architecture, implementation, validation, and scale plan"
    document.core_properties.author = "Industrial Data Catalogs program"
    document.core_properties.created = datetime(2026, 8, 7, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 8, 7, tzinfo=UTC)
    document.core_properties.revision = 1
    document.core_properties.keywords = "NVIDIA, Nemotron, industrial catalogs, VLM, RAG, extraction, architecture"
    document.core_properties.comments = "Generated from the implemented local project and verified NVIDIA server run evidence."

    add_cover(document)
    body_section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    body_section.different_first_page_header_footer = False
    set_page_number_start(body_section, 1)
    configure_header_footer(body_section)

    add_report_map(document)

    add_label(document, "Decision brief")
    add_title(document, "1. Executive Summary", 1)
    add_body(document, "The NVIDIA Industrial Catalog Extraction Pipeline is an evidence-preserving system for converting heterogeneous industrial PDFs into grounded product records and a citation-bearing retrieval knowledge base. Its defining rule is that model output does not become canonical merely because inference succeeded: every field must resolve to page evidence, survive deterministic source guards, and pass typed validation before persistence.")
    add_body(document, "The system uses three NVIDIA model services on dedicated A100 GPUs. NVIDIA Nemotron Parse v1.2 reads complex document pages; Llama 3.1 Nemotron Nano 8B v1 converts evidence into schema-constrained product JSON; and Nemotron 3 Embed 8B BF16 creates retrieval vectors. Native PDF extraction handles clean digital pages without paying VLM cost, while an optional OCR verification hook can add independent identifier evidence when deployed.")
    add_title(document, "What is implemented and proven", 2)
    add_bullet(document, "End-to-end code, deployment scripts, service health checks, deterministic routing, extraction, guarded parsing, SQLite persistence, two-tier RAG, monitoring, checkpoints, and audit artifacts are implemented.")
    add_bullet(document, "The verified server exposes eight NVIDIA A100-SXM4 80 GB GPUs, approximately 2 TB of system memory, 255 logical CPUs, and roughly 17.4 TB free on the data volume at inspection time.")
    add_bullet(document, "A canonical transfer containing 220 PDFs and 17,837 pages was hash-verified on the server. Vendor catalogs and the broader reference library are treated as separate campaigns because their semantics differ.")
    add_bullet(document, "A 20-document production validation traversed every stage; a hardened two-document replay then resolved both observed extraction failures while keeping unsupported products out of the catalog.")
    add_bullet(document, "Local and remote verification completed with more than 140 tests passing and Ruff reporting no lint findings.")
    add_title(document, "What is not yet complete", 2)
    add_body(document, "The complete 17,837-page campaign has not been executed. The current benchmark wrapper is intentionally bounded and the runner is sequential. Whole-document knowledge-base replacement also prevents safe page-window sharding today. A manifest-pinned campaign runner, isolated work-unit outputs, and deterministic merge must be implemented before full-corpus completion can be claimed.")
    add_callout(document, "Recommended decision", "Approve the manifest/sharding/merge work as the next production milestone, then run a two-worker vendor pilot. Expand to a third worker only when aggregate throughput improves without rising queue time, invalid JSON, timeouts, or GPU pressure.", kind="success")

    add_label(document, "Purpose and boundaries")
    add_title(document, "2. Objectives, Scope, and Design Principles", 1)
    add_title(document, "2.1 Business objective", 2)
    add_body(document, "Industrial catalogs contain high-value product identity and engineering data, but layouts vary by vendor, edition, page type, and scan quality. The pipeline’s purpose is to turn those files into searchable, source-grounded records that can support product discovery, specification lookup, comparison, and downstream retrieval-augmented generation.")
    add_table(
        document,
        ("Target output", "Examples", "Admission requirement"),
        (
            ("Product identity", "Manufacturer, product/part name, part number, series, family", "Exact or source-justified evidence and deterministic manufacturer/cardinality rules"),
            ("Specifications", "Voltage, dimensions, material, rating, tolerance, capacity", "Typed value plus source evidence; no unsupported normalization"),
            ("Operating details", "Conditions, compatibility, compliance, application notes", "Validated field with document/page/element citation"),
            ("Other details", "Description, ordering context, features, warnings", "Schema-valid content traceable to extracted evidence"),
            ("Retrieval evidence", "Page element, exact excerpt, bbox, score", "Indexed independently of canonical product admission"),
        ),
        (1750, 3310, 4300),
    )
    add_title(document, "2.2 Architectural principles", 2)
    principles = [
        ("Evidence first", "Keep exact source text, document identity, page, bounding box, extraction method, provider payload, and model metadata."),
        ("Deterministic around probabilistic", "Use explicit thresholds, stable identifiers, strict schemas, source guards, typed validation, and idempotent replacement."),
        ("Fail closed", "Do not convert blank/unsupported model output or unresolved nonblank pages into canonical facts."),
        ("Selective acceleration", "Use native PDF evidence when it is good; reserve the document VLM for scans, tables, columns, and visually complex layouts."),
        ("Two-tier retrieval", "Keep raw evidence searchable even when a candidate does not qualify for the canonical product catalog."),
        ("Operational boundedness", "Limit candidate pages, wall time, GPUs, and output directories until campaign-scale correctness is implemented."),
    ]
    add_table(document, ("Principle", "How the implementation expresses it"), principles, (1900, 7460), font_size=9.1)

    add_label(document, "Verified platform")
    add_title(document, "3. NVIDIA Infrastructure, Model Stack, and Corpus", 1)
    add_title(document, "3.1 Server capacity", 2)
    add_table(
        document,
        ("Resource", "Verified capacity", "Pipeline use"),
        (
            ("GPU", "8 × NVIDIA A100-SXM4 80 GB (640 GB aggregate VRAM)", "Three isolated inference services; two GPUs reserved; three free for measured expansion"),
            ("System memory", "Approximately 2 TB RAM", "Large page buffers, concurrent workers, model/service overhead, and merge operations"),
            ("CPU", "255 logical CPUs", "PDF parsing, image rendering, orchestration, validation, and SQLite work"),
            ("Data volume", "About 20.4 TB total / 17.4 TB free at inspection", "Corpus, checkpoints, run artifacts, databases, vectors, and published releases"),
        ),
        (1900, 3000, 4460),
    )
    add_title(document, "3.2 NVIDIA model responsibilities", 2)
    add_table(
        document,
        ("Stage", "Model / endpoint", "GPU", "Role and operating contract"),
        (
            ("Document VLM", "nvidia/NVIDIA-Nemotron-Parse-v1.2\n127.0.0.1:8001", "1", "Visual document parsing: reading order, text, tables, element classes, and normalized bounding boxes; max sequences 8"),
            ("Structured LLM", "nvidia/Llama-3.1-Nemotron-Nano-8B-v1\n127.0.0.1:8003", "3", "Guided JSON for product identity/specifications; source hints and post-response guard; max sequences 16"),
            ("Embeddings", "nvidia/Nemotron-3-Embed-8B-BF16\n127.0.0.1:8004", "5", "Passage/query embeddings for exact cosine retrieval; production output documented as 4,096 dimensions; max sequences 16"),
        ),
        (1550, 3150, 650, 4010),
        font_size=8.4,
    )
    add_body(document, "The structured LLM runs through vLLM 0.20.0 with the xgrammar structured-output backend and disabled arbitrary whitespace. GPU 0 and GPU 2 belong to pre-existing workloads and are outside project control. GPUs 4, 6, and 7 remain unassigned until a measured capacity plan authorizes them.")
    add_title(document, "3.3 Corpus inventory and campaign separation", 2)
    add_table(
        document,
        ("Corpus", "PDFs", "Pages", "Intended extraction profile"),
        (
            ("Vendor product catalogs", "110", "6,069", "Manufacturer, part identity, specifications, operating details"),
            ("Industrial reference library", "110", "11,768", "Reference-document extraction with its own quality policy"),
            ("Combined transfer", "220", "17,837", "Two pinned campaigns; optionally merged into one retrieval release"),
        ),
        (2600, 900, 1100, 4760),
    )
    add_callout(document, "Coverage caveat", "Nineteen vendor PDFs exceed 50 pages and contain 5,247 pages. A 50-page-per-document cap covers only 1,772 of 6,069 vendor pages (29.2%); raising only the document limit cannot create full coverage.", kind="warning")
    add_body(document, "The research corpus was transferred through a hash-verified 918,048,768-byte archive. Neither that archive nor the source PDFs are distributed with the public repository. The exact deployment code and data identities were retained in the private operational record.")

    add_label(document, "System structure")
    add_title(document, "4. End-to-End Architecture", 1)
    add_figure(
        document,
        architecture_path,
        "Figure 1. End-to-end evidence and knowledge-base architecture.",
        "Flow from PDF intake through deterministic routing, NVIDIA Nemotron Parse, evidence creation, Nemotron Nano structured parsing, source guard, validation, storage, Nemotron Embed, and two-tier retrieval.",
    )
    add_title(document, "4.1 Component responsibilities", 2)
    add_table(
        document,
        ("Component", "Input", "Output", "Control boundary"),
        (
            ("Document intake", "PDF file", "Document SHA, page iterator, metadata", "Content-addressed identity; lazy page rendering"),
            ("Page router", "Native signals + layout inventory", "native_text / document_vlm / OCR flags + reasons", "Fixed thresholds and ordered reason codes"),
            ("Nemotron Parse", "Rendered page image + task prompt", "Text/table/element evidence", "Bounded retries; semantic response validation"),
            ("Evidence layer", "Native/VLM/OCR elements", "Canonical elements with IDs and citations", "No model-generated field is evidence by itself"),
            ("Nemotron Nano", "Bounded evidence chunk + guided schema", "Product candidates", "Manufacturer and cardinality hints constrain output"),
            ("Guard + validator", "Candidates + canonical evidence", "Accepted, rejected, or invalid records", "Cross-document and unsupported claims fail closed"),
            ("Catalog + RAG", "Accepted products + all page elements", "SQLite/JSONL + vectors/citations", "Canonical tier separated from raw evidence tier"),
        ),
        (1550, 2170, 2680, 2960),
        font_size=8.3,
    )

    add_label(document, "Extraction mechanics")
    add_title(document, "5. How the Pipeline Works", 1)
    add_title(document, "5.1 Document identity and page intake", 2)
    add_number(document, "Hash the PDF and assign the stable document ID sha256:<file hash>.")
    add_number(document, "Stream pages in one-based order with PyMuPDF; do not render every page eagerly.")
    add_number(document, "Inventory native text/blocks, page dimensions, tables, multi-column structure, image coverage, drawings, embedded images, annotations, and links.")
    add_number(document, "Render PNG bytes only when the selected VLM/OCR path requires pixels.")
    add_body(document, "This division keeps large digital catalogs inexpensive while retaining the source-object detail needed to distinguish truly blank pages from nontext visual content.")

    add_title(document, "5.2 Deterministic page routing", 2)
    add_table(
        document,
        ("Signal", "Default threshold / trigger", "Interpretation"),
        (
            ("Non-whitespace characters", "≥ 80", "Minimum native text volume"),
            ("Words", "≥ 12", "Rejects thin headers and fragments"),
            ("Printable ratio", "≥ 0.95", "Avoids corrupt encoding"),
            ("Alphanumeric ratio", "≥ 0.35", "Requires meaningful text content"),
            ("Replacement-character ratio", "≤ 0.01", "Limits decoding damage"),
            ("Layout trigger", "Table, multi-column, or image coverage ≥ 0.45", "Routes page to the document VLM despite usable native text"),
        ),
        (2400, 2300, 4660),
    )
    add_body(document, "The routing decision stores every measured signal and an ordered reason-code list. A page can therefore be reproduced and explained without asking the model why it was selected.")

    add_title(document, "5.3 NVIDIA Nemotron Parse extraction", 2)
    add_body(document, "VLM-routed pages are rendered and submitted as base64 image URLs to the local OpenAI-compatible endpoint. Parse v1.2 receives the task prompt shown below; sampling is deterministic (temperature 0, top_k 1) and output length is bounded to 8,192 tokens.")
    add_code_block(document, [
        "</s><s><predict_bbox><predict_classes><output_markdown>",
        "<predict_no_text_in_pic>",
        "temperature=0  top_k=1  max_tokens=8192  repetition_penalty=1.1",
    ])
    add_body(document, "Normalization accepts JSON element payloads, Nemotron’s tagged class/text/bounding-box format, or a final document-text element when a supported response cannot be split further. Tagged Parse bounding boxes are explicitly labeled as normalized coordinates; native blocks retain PyMuPDF page coordinates.")

    add_title(document, "5.4 Optional OCR verification", 2)
    add_body(document, "OCR is implemented as an independent verifier, not a destructive replacement. When an endpoint is configured, VLM pages can be OCR-checked and native pages containing identifier-like tokens can receive an additional evidence channel. The current production example does not deploy an OCR model, so the hook is present but inactive.")

    add_title(document, "5.5 Grounded evidence elements", 2)
    add_table(
        document,
        ("Field group", "Preserved content"),
        (
            ("Source identity", "Document ID, source path, one-based page, sequence number"),
            ("Evidence", "Exact text, element type, bounding box, confidence"),
            ("Provenance", "Extraction method, model identity, provider payload, metadata"),
            ("Determinism", "Element ID hashed from document/page/sequence/type/text/method/bbox/model"),
            ("Audit", "Route, reason codes, attempts, failures, fallback/blank events"),
        ),
        (2350, 7010),
    )

    add_label(document, "Structured intelligence")
    add_title(document, "6. Product Parsing, Source Guard, and Validation", 1)
    add_title(document, "6.1 Pre-model source hints", 2)
    add_body(document, "Before Nemotron Nano is called, deterministic code examines the evidence for manufacturer candidates and product cardinality. Manufacturer hints come from legal names, labels, copyright lines, and brand evidence. Cardinality comes from explicit orderable rows rather than the LLM’s preference to split or combine products.")
    add_table(
        document,
        ("Source condition", "Guided-schema effect", "Post-response enforcement"),
        (
            ("Explicit manufacturer evidence", "Manufacturer and evidence IDs constrained to source-backed values", "Unrecognized manufacturer/evidence pairs are rejected"),
            ("Ordinary prose", "At most one product candidate", "Extra product objects are guard-rejected"),
            ("Distinct orderable table rows", "One identity-only product per exact identifier", "Identifiers and row cardinality must match evidence"),
            ("No justified part-number structure", "Part-number extraction disabled", "Invented or inferred numbers are rejected"),
        ),
        (2400, 3370, 3590),
    )
    add_title(document, "6.2 Guided JSON and bounded regeneration", 2)
    add_body(document, "Nemotron Nano receives a strict JSON Schema through OpenAI response_format. Malformed or truncated JSON can trigger one bounded regeneration request. Attempts, finish reasons, provider responses, decoded content source, and error classification remain available for audit.")
    add_table(
        document,
        ("Structured wire contract", "Implemented bound / behavior"),
        (
            ("Top level", "Exactly one object with a required products array; unknown properties rejected"),
            ("Product cardinality", "Maximum 32 per parse batch before stricter source-derived bounds"),
            ("Required identity", "Manufacturer and part name; part number optional and disabled on single-family pages"),
            ("Specifications", "Name and value required; optional unit; maximum 48"),
            ("Other details", "Fixed name/value object array; maximum 24"),
            ("Evidence references", "One to eight supplied element IDs per emitted value"),
            ("Raw string", "Nonblank and at most 2,048 characters"),
        ),
        (3000, 6360),
        font_size=8.7,
    )
    add_body(document, "The wire object is intentionally smaller than the canonical product model. Trusted application code adds source identity, canonical evidence, normalized values, confidence, IDs, extraction metadata, and review status. Canonical materials, certifications, operating conditions, compatible parts, and aliases exist for future enrichment but are not currently emitted by the guided prompt.")
    add_title(document, "6.3 Versioned contracts", 2)
    add_table(
        document,
        ("Contract", "Version"),
        (
            ("Extraction pipeline", "1.1"),
            ("Canonical ProductRecord / ProductBatch", "1.0"),
            ("Catalog SQLite schema", "2"),
            ("Knowledge-base schema", "1.1"),
            ("Source guard", "source-guard-v3"),
            ("Identity-row repair", "identity-row-repair-v2"),
        ),
        (6500, 2860),
        font_size=8.9,
    )
    add_title(document, "6.4 Evidence hydration and typed validation", 2)
    add_body(document, "Compact evidence references returned by the LLM are replaced with canonical extraction elements. Validation rejects unknown element IDs, cross-document evidence, mismatched claimed text or page values, missing required evidence, malformed fields, and unsupported normalizations. Valid siblings may still be stored when another candidate in the same ingestion batch is rejected.")
    add_callout(document, "Canonical admission rule", "A plausible value without traceable evidence is not a product fact. It may remain searchable in raw page evidence, but it cannot enter the canonical product catalog.", kind="success")
    add_table(
        document,
        ("Validation outcome", "Representative conditions", "Result"),
        (
            ("Always error", "Schema failure; unknown/cross-document evidence; no part name or number; multiline/replacement-character part number", "Candidate rejected"),
            ("Warning by default", "Missing manufacturer, part number, specifications, or field evidence", "Accepted as needs_review; becomes error in strict mode"),
            ("Warning", "Confidence below 0.65; duplicate/conflicting specifications; blank value; evidence-claim mismatch", "Accepted with audit warning when no error exists"),
            ("Accepted", "Canonical model exists and no error-level issue remains", "Persisted and eligible for canonical RAG"),
            ("Review document", "At least one accepted product carries a warning or quality issue", "Terminal review status, kept separate from failure"),
        ),
        (1850, 5000, 2510),
        font_size=8.35,
    )

    add_label(document, "Resilience and safety")
    add_title(document, "7. Retry, Blank-Page, Fallback, and Failure Policies", 1)
    add_figure(
        document,
        resilience_path,
        "Figure 2. Page extraction decision and bounded recovery policy.",
        "Decision tree for certified blank pages, deterministic routing, Nemotron Parse semantic retries, narrowly permitted native fallback, evidence normalization, and fail-closed unresolved pages.",
    )
    add_title(document, "7.1 Transport and semantic retry", 2)
    add_body(document, "The shared client uses deterministic exponential retry without jitter. It retries transport failures and HTTP 408, 409, 425, 429, 500, 502, 503, and 504. The Parse configuration specifies seven retries in addition to the first request, permitting eight total attempts.")
    add_body(document, "A 2xx response is not automatically accepted. The client searches supported assistant payload locations—message content, parsed objects, content blocks, tool/function arguments, completion text, reasoning-content fallbacks, and Responses-style output. Blank or unsupported successful responses are audited and retried.")
    add_title(document, "7.2 Certified blank pages", 2)
    add_body(document, "Blankness is positively certified only when the source-object inventory is complete and native text, native blocks, tables, image coverage, preloaded image bytes, drawings, images, annotations, links, and multi-column structure all indicate zero/false. Missing or invalid signals never count as blank. Certified blank pages emit no element, write one stable audit, and permit the structured LLM to be skipped.")
    add_title(document, "7.3 Native fallback after semantic exhaustion", 2)
    add_body(document, "Fallback is permitted only for status 200 with failure kind unsupported_assistant_content, at least one successful-but-invalid HTTP response, and independently quality-approved native text. The original route stays document_vlm and the full failure record is retained. Transport errors, server failures, mixed error sequences, and nonblank pages without meaningful native text fail closed.")
    add_table(
        document,
        ("Observed condition", "Terminal action", "Why"),
        (
            ("Exact zero source-object inventory", "Certified empty; skip structured parse", "Avoids unnecessary inference without hiding content"),
            ("HTTP 200, supported content", "Normalize to evidence", "Semantic contract satisfied"),
            ("HTTP 200, blank/unsupported content", "Retry; retain raw response", "Transport success did not produce evidence"),
            ("Semantic exhaustion + good native text", "Audited native fallback", "Independent evidence is available"),
            ("Nonblank page without usable native evidence", "Fail closed", "No safe source-grounded substitute"),
            ("Transport/server/mixed retry exhaustion", "Fail closed", "Does not meet the narrow fallback predicate"),
        ),
        (3090, 2600, 3670),
        font_size=8.6,
    )

    add_label(document, "Canonical data and retrieval")
    add_title(document, "8. Storage, Data Model, and Two-Tier RAG", 1)
    add_title(document, "8.1 Canonical storage model", 2)
    add_table(
        document,
        ("Entity / artifact", "Purpose", "Lineage or integrity behavior"),
        (
            ("documents", "Document status and identity", "Content-addressed by SHA; review remains distinct from failure"),
            ("extracted elements", "Page-grounded evidence", "Stable element IDs; exact text, bbox, method, model, audit"),
            ("ingestion batches", "One structured parse/validation decision unit", "Stores raw extraction, raw LLM output, accepted/rejected counts"),
            ("products", "Canonical product identity and details", "Only source-guarded and validated records"),
            ("specifications / conditions", "Typed engineering facts", "Evidence links retained at field level"),
            ("knowledge_chunks", "Retrieval text and embedding", "Deterministic content-derived chunk IDs and owner-scoped replacement"),
            ("citations", "Exact retrieval provenance", "Document/path/page/element/bbox/excerpt/confidence/field paths"),
        ),
        (2000, 3280, 4080),
        font_size=8.6,
    )
    add_title(document, "8.2 Two retrieval tiers", 2)
    add_table(
        document,
        ("Tier", "Chunking", "Admission rule", "Primary use"),
        (
            ("Raw page evidence", "One chunk per nonblank extracted page element", "Indexed before product admission", "Find exact page facts, rejected candidates, and unresolved content"),
            ("Canonical product", "Identity plus one chunk per specification/operating condition", "Only validated persisted products", "Reliable product search and grounded downstream RAG"),
        ),
        (1750, 2900, 2360, 2350),
        font_size=8.6,
    )
    add_body(document, "Indexed passages are prefixed with passage: and queries with query:. The current store keeps float vectors as JSON in SQLite and computes exact cosine scores over all compatible vectors. Text-only mode performs escaped case-insensitive substring matching with deterministic structural ordering.")
    add_body(document, "The library API can restrict retrieval by chunk tier, but the current CLI searches product and raw page-evidence tiers together. A user-facing tier selector and tier-aware ranking policy should be added before the interface is presented as a production search product.")
    add_callout(document, "Retrieval qualification", "The production embedding model is documented as 4,096-dimensional, but current runtime validation checks only nonempty, finite, batch-consistent vectors. Model and exact dimension should become enforced release invariants before a large merged knowledge base is published.", kind="warning")
    add_title(document, "8.3 Determinism and replacement", 2)
    add_body(document, "Stable JSON serialization and SHA-256 over text, citations, metadata, type, ordinal, and owner create content-derived chunk IDs. Embedding completes before replacement. Product, page, and document scopes are deleted and reinserted inside SQLite transactions, making retries idempotent at each owner scope. A multi-product indexing call, however, commits once per product rather than as one collection-wide transaction.")

    add_label(document, "Recovery and evidence chain")
    add_title(document, "9. Checkpointing, Auditability, and Output Artifacts", 1)
    add_title(document, "9.1 Page checkpoints", 2)
    add_body(document, "Checkpoint fingerprints include pipeline version 1.1, document/page identity, native text and block structure, dimensions, source-object metadata, preloaded image hash when present, complete routing decision, and credential-free Parse/OCR client configuration. Checkpoints are written through a temporary file, flushed, fsync’d, and atomically replaced. Restore requires matching fingerprint, document ID, and page.")
    add_callout(document, "Known invalidation gap", "Render DPI/profile is not an explicit fingerprint field for lazy page images. The source-PDF hash protects document identity, but a rendering-profile change should directly invalidate VLM page checkpoints in the campaign runner.", kind="warning")
    add_title(document, "9.2 Run artifacts", 2)
    add_table(
        document,
        ("Artifact", "Contents", "Use"),
        (
            ("summary.json", "Run identity, elapsed time, document/page/product/chunk counts", "Operator reconciliation and release gating"),
            ("documents.jsonl", "Per-document terminal status and failure/review context", "Audit, retry selection, and status reporting"),
            ("extraction/", "Page evidence, routing, provider payloads, checkpoints", "Reproduction and page-level investigation"),
            ("structured/", "Guided JSON attempts and failure payloads", "Decode/finish-reason analysis and bounded replay"),
            ("catalog.sqlite3", "Canonical records, evidence, ingestion batches, rejections", "Transactional product system of record"),
            ("knowledge_base.sqlite3", "Chunks, vectors, citations", "Citation-bearing retrieval"),
            ("monitor.latest.json / history", "GPU, endpoint, data volume, job state", "Live watch and post-run proof"),
        ),
        (2330, 3840, 3190),
        font_size=8.6,
    )
    add_title(document, "9.3 Audit chain", 2)
    add_body(document, "A reviewer can move from a retrieved chunk to its citation, from citation to exact element and page, from element to routing/extraction method, and from model-derived product field back through source guard, validation, provider response, and ingestion batch. Rejected products remain journaled; they are not silently discarded.")

    add_label(document, "Deployment and operations")
    add_title(document, "10. NVIDIA Service Topology and Run Operations", 1)
    add_figure(
        document,
        topology_path,
        "Figure 3. Verified GPU service assignment and project guardrails.",
        "Server topology showing the runner and monitoring layer, Nemotron Parse on GPU 1, Nemotron Nano on GPU 3, Nemotron Embed on GPU 5, reserved GPUs 0 and 2, and free GPUs 4, 6, and 7.",
    )
    add_title(document, "10.1 Deployment locations", 2)
    add_table(
        document,
        ("Location", "Path", "Purpose"),
        (
            ("Repository checkout", "${PROJECT_ROOT}", "Code, tests, configuration, scripts, and this report"),
            ("NVIDIA server project", "${SERVER_PROJECT_ROOT}", "Deployed code on the NVIDIA server"),
            ("Remote data root", "/data/industrial-data-corps", "Corpus, manifests, environments, runs, status, and future releases"),
        ),
        (1700, 5050, 2610),
        font_size=8.4,
    )
    add_title(document, "10.2 Service lifecycle", 2)
    add_code_block(document, [
        "bash scripts/bootstrap_remote.sh",
        "bash scripts/start_model_services.sh --with-embedding",
        "curl -fsS http://127.0.0.1:8001/health",
        "curl -fsS http://127.0.0.1:8003/health",
        "curl -fsS http://127.0.0.1:8004/health",
        "bash scripts/stop_model_services.sh   # scoped project shutdown only",
    ])
    add_body(document, "The launcher verifies PID, model, port, physical GPU, health, and advertised model identity before adopting a live service. Operational policy prohibits broad process kills or GPU resets because GPUs 0 and 2 carry unrelated workloads.")
    add_title(document, "10.3 Bounded validation envelope", 2)
    add_body(document, "scripts/run_benchmark.sh defaults to 20 documents, 12 pages per document, at most 240 candidate pages, and a two-hour wall timeout. It enforces 1–100 documents, 1–50 pages per document, a 500-page maximum per invocation, 60–86,400 seconds wall time, approved physical GPU visibility, and one unique run directory.")
    add_code_block(document, [
        "INDUSTRIAL_BENCHMARK_MAX_DOCUMENTS=20 \\",
        "INDUSTRIAL_BENCHMARK_MAX_PAGES_PER_DOCUMENT=12 \\",
        "INDUSTRIAL_BENCHMARK_TIMEOUT_SECONDS=7200 \\",
        "bash scripts/run_benchmark.sh --live --pages-per-parse 1",
    ])
    add_title(document, "10.4 Monitoring and stop conditions", 2)
    add_bullet(document, "Monitor health and model identity on all three endpoints; confirm only GPUs 1, 3, and 5 are owned by the project services.")
    add_bullet(document, "Stop or pause on endpoint degradation, repeated invalid/truncated JSON, increasing HTTP retries/timeouts, OOM or site temperature threshold, reserved-GPU ownership, SQLite integrity/locking errors, or a data volume above 90% use.")
    add_bullet(document, "Keep review separate from failure. A review record contains accepted products with warnings; it is not an execution failure.")
    add_bullet(document, "Never share an output directory between processes. JSONL writers are process-local and document-level evidence replacement can overwrite concurrent work.")

    add_label(document, "Measured proof")
    add_title(document, "11. Validation Results and Incident Replay", 1)
    add_title(document, "11.1 Production validation batch", 2)
    add_body(document, "Run 20260807T064527Z-971750 exercised the complete pipeline on a bounded selection. It demonstrates integration and provides capacity evidence; it does not represent corpus completion.")
    add_table(
        document,
        ("Metric", "Value", "Interpretation"),
        (
            ("Documents attempted", "20", "Bounded validation selection"),
            ("Review", "18", "Accepted with quality warnings; not failed"),
            ("Failed", "2", "Both later resolved by hardened replay"),
            ("Pages counted", "105", "Pages in successful/review documents"),
            ("Extracted elements", "1,622", "Grounded page evidence"),
            ("Products parsed", "73", "Structured candidates before all controls"),
            ("Guard-rejected", "41", "Unsupported source/cardinality claims blocked"),
            ("Products stored", "31", "Canonical accepted records"),
            ("Invalid", "1", "Failed typed/evidence validation"),
            ("Knowledge chunks", "1,655", "Raw + canonical retrieval material"),
            ("Elapsed", "828.999771 s", "Includes failed-document time"),
        ),
        (2870, 1600, 4890),
        font_size=8.6,
    )
    add_title(document, "11.2 Hardened replay of the two failures", 2)
    add_body(document, "Run 20260807T073659Z-1185130 processed one vendor catalog and one reference document after resilience hardening. It completed 24 pages in 85.468255 seconds with two review documents, zero failures, 531 extracted elements, and 531 page-evidence chunks/citations.")
    add_table(
        document,
        ("Case", "Observed condition", "Correct terminal behavior"),
        (
            ("Vendor catalog page 6", "Eight successful HTTP responses contained no supported assistant content", "Used 25 quality-approved native blocks; retained original VLM route and every response/fallback audit"),
            ("Reference document page 2", "Complete page inventory proved no text, blocks, tables, images, drawings, annotations, or links", "Certified blank; emitted no element; skipped structured LLM; created no ingestion batch"),
            ("Replay candidate", "One structured product candidate did not satisfy source guard", "Rejected deterministically; zero unsupported products entered the catalog"),
        ),
        (1950, 3650, 3760),
        font_size=8.7,
    )
    add_body(document, "The replay produced 23 ingestion batches for 24 page chunks, exactly matching the single certified blank skip. Both SQLite databases reported integrity_check=ok and no foreign-key violations. A live NVIDIA-vector query returned grounded vendor-catalog evidence, including the page-6 native fallback citation. Final endpoint checks returned HTTP 200.")
    add_callout(document, "Quality signal", "Rejecting 41 of 73 validation candidates and the only replay candidate is evidence that the source guard is active. High storage count is not the optimization target; source-grounded precision is.", kind="success")

    add_label(document, "Capacity interpretation")
    add_title(document, "12. Throughput and Capacity Planning", 1)
    add_body(document, "The validation batch observed 7.90 seconds per counted page, 7.6 pages per minute, or roughly 456 pages per hour. Because failed-document time is included while failed pages are excluded from the denominator, this is a planning observation—not a clean page-only benchmark, completion claim, or service-level objective.")
    add_table(
        document,
        ("Scope", "Pages", "Arithmetic single-worker baseline", "Operational planning range"),
        (
            ("Vendor catalogs", "6,069", "13.3 hours", "16.6–20 hours"),
            ("Reference library", "11,768", "25.8 hours", "Pilot separately before commitment"),
            ("Combined", "17,837", "39.1 hours", "49–59 hours before measured concurrency gains"),
        ),
        (2250, 1200, 2830, 3080),
        font_size=8.9,
    )
    add_title(document, "12.1 Bottlenecks to measure", 2)
    add_bullet(document, "Nemotron Parse latency and queue depth on image-heavy/table pages.")
    add_bullet(document, "Structured JSON completion length, regeneration rate, and source-guard rejection distribution.")
    add_bullet(document, "Embedding batch latency and SQLite vector serialization size.")
    add_bullet(document, "CPU rendering/parse time, data-volume write pressure, checkpoint hit rate, and merge cost.")
    add_bullet(document, "Aggregate pages/hour and p95 stage latency as worker count moves from one to two, then conditionally to three.")

    add_label(document, "Risk governance")
    add_title(document, "13. Limitations and Risk Register", 1)
    add_table(
        document,
        ("Risk / limitation", "Impact", "Current control", "Required next control"),
        (
            ("Runner is sequential; worker settings are unused", "No automatic scale-out", "Bounded runs and known capacity baseline", "External manifest-assigned shard workers"),
            ("No manifest/page-range campaign ledger", "Cannot prove full coverage or safe retry", "Inventory files and unique bounded runs", "Immutable campaign manifest + stable work-unit state"),
            ("Whole-document evidence replacement", "Page-window jobs can erase sibling evidence", "Do not shard one document today", "Shard-local DBs + deterministic merge"),
            ("OCR endpoint not deployed", "Some scan/identifier cases depend on VLM only", "Fail closed; OCR hook exists", "Deploy and benchmark a dedicated NVIDIA OCR verifier"),
            ("Exact cosine full scan over JSON vectors", "Query latency grows linearly", "Acceptable for bounded KB", "sqlite-vec or external vector index; optional reranker"),
            ("Embedding dimension/model not enforced", "Mixed vectors may silently skip or compare incorrectly", "Batch-consistency checks", "Release-wide model and exact 4,096-D invariant"),
            ("Citation list not schema-required for every chunk", "Direct product chunks could lack evidence", "Pipeline-created page chunks always cite", "Require ≥1 citation for canonical chunks"),
            ("Render profile absent from fingerprint", "Changed page rendering may reuse stale VLM checkpoint", "Source PDF hash", "Fingerprint DPI/profile and prompt version explicitly"),
            ("No general stored-response replay command", "Investigation may repeat LLM/embedding work", "Retain raw artifacts; rerun exact doc in new directory", "Replay tool that consumes stored provider responses"),
            ("Field-to-excerpt semantic proof is partial", "A cited element may not fully entail an emitted value outside guarded identity cases", "Unknown/cross-document evidence fails; manufacturer/row checks are strict", "Exact span/value checks for identity and specification fields"),
            ("JSONL and SQLite are not one transaction", "Crash timing can duplicate an event", "Written flags suppress ordinary retries", "SQLite outbox or deterministic release-time JSONL generation"),
            ("Reprocessing lacks source-scope tombstones", "Older products can survive a parse that emits fewer products", "Owner-scoped product and chunk replacement", "Versioned source reconciliation and supersession policy"),
            ("Product IDs contain model-output ordinal", "Reordering equivalent products changes identifiers", "Canonical JSON participates in deterministic ID", "Prefer stable source identity keys when available"),
            ("Model revision is not in batch identity", "A model change may reuse a logical batch ID when parsed data matches", "Pipeline/schema hashes are included", "Add stable configuration and model-revision hash"),
            ("CLI mixes raw and canonical RAG tiers", "Users may conflate evidence with validated product data", "Chunk kind retained; API can filter", "Expose tier filter and explicit tier labeling/ranking"),
        ),
        (2300, 2050, 2380, 2630),
        font_size=7.9,
    )
    add_callout(document, "Primary production risk", "The principal blocker is correctness of campaign partitioning and merge—not GPU capacity. The server has ample compute, but uncontrolled parallelism could create gaps, overlaps, overwritten evidence, or a release that cannot prove page coverage.", kind="danger")

    add_label(document, "Production roadmap")
    add_title(document, "14. Full-Corpus Scale Plan", 1)
    add_title(document, "14.1 Immutable campaign manifests", 2)
    add_body(document, "Create separate canonical manifests for the vendor corpus and reference library. Each document entry must include corpus-relative path, SHA-256, byte size, and total pages. Canonically hash the manifest and copy that identity into every unit status, attempt summary, and merged release. Reject escaped paths, symlinks, altered hashes, altered page counts, or ambiguous duplicate-content policy before inference.")
    add_title(document, "14.2 Stable page-window work units", 2)
    add_body(document, "Split PDFs into nonoverlapping one-based inclusive windows no larger than 50 pages. Persist identities such as sha-prefix-p000001-000050; never regenerate assignments during retry. Each unit writes to an isolated attempt directory and a stable checkpoint directory. Completion markers record source/campaign/config/prompt/model identities, page range, attempt, timestamps, metrics, and output hashes.")
    add_title(document, "14.3 External workers and measured promotion", 2)
    add_table(
        document,
        ("Campaign", "50-page work units", "Initial shard layout", "Worker policy"),
        (
            ("Vendor", "206", "~24 shards averaging ~253 pages", "Start with 2 workers; promote to 3 only after a stable concurrency trial"),
            ("Reference library", "302", "~48 similar shards after a profile-specific pilot", "Tune routing/quality separately; do not inherit vendor assumptions blindly"),
        ),
        (1750, 1770, 2850, 2990),
        font_size=8.7,
    )
    add_title(document, "14.4 Deterministic merge and release gates", 2)
    for text in (
        "Accept only completed/review work units whose campaign, source, pipeline, prompt, configuration, and model identities match the release.",
        "Create a new temporary release; import catalog rows in foreign-key order and vectors without re-embedding. Identical duplicate keys may deduplicate; conflicting payloads fail closed.",
        "Require every manifest page exactly once with no gaps/overlaps, no pending/running/unresolved failed units, and separate review counts.",
        "Require SQLite integrity_check=ok, an empty foreign_key_check, reconciled document/product/chunk/citation counts, one embedding model/dimension, release summary, and checksums.",
        "Publish atomically only after all gates pass; keep shard/attempt artifacts as the long-term audit source.",
    ):
        add_bullet(document, text)
    add_title(document, "14.5 Delivery sequence", 2)
    roadmap = [
        ("Phase 1", "Campaign correctness", "Manifest builder, page windows, unit ledger, isolated outputs, deterministic merge, release verifier", "All new unit/merge tests pass; synthetic gaps/conflicts fail closed"),
        ("Phase 2", "Vendor pilot", "Two external workers over representative short/long/table/scan documents", "Improved aggregate throughput; no queue/OOM/JSON regression; citation QA passes"),
        ("Phase 3", "Vendor campaign", "206 work units, ~24 shards, monitored execution and merge", "6,069 pages exactly once; zero unresolved failures; published verified release"),
        ("Phase 4", "Library pilot/campaign", "Separate extraction schema/quality profile followed by ~302 units", "11,768 pages exactly once under library-specific acceptance policy"),
        ("Phase 5", "Combined retrieval release", "Conflict-checked merge, vector-index upgrade, retrieval evaluation", "Unified citation-bearing KB with release manifest and model/dimension invariant"),
    ]
    add_table(document, ("Phase", "Objective", "Core work", "Exit gate"), roadmap, (1000, 1800, 3520, 3040), font_size=8.1)

    add_label(document, "Implementation guide")
    add_title(document, "15. Key Files and Operator Entry Points", 1)
    add_table(
        document,
        ("File", "Responsibility"),
        (
            ("src/industrial_catalog/routing.py", "Native-quality/layout signals and deterministic route selection"),
            ("src/industrial_catalog/extraction.py", "PDF source, page inventory, Parse/OCR orchestration, normalization, blank/fallback policy, checkpoints"),
            ("src/industrial_catalog/nvidia_clients.py", "OpenAI-compatible NVIDIA requests, retries, content candidates, Parse/OCR/guided JSON clients"),
            ("src/industrial_catalog/validation.py", "Canonical evidence hydration, field/product validation, quality issues"),
            ("src/industrial_catalog/storage.py", "Catalog schema, ingestion audit, accepted/rejected persistence"),
            ("src/industrial_catalog/knowledge_base.py", "Chunk/citation models, embeddings, deterministic indexing, text and cosine retrieval"),
            ("src/industrial_catalog/runner.py", "End-to-end bounded execution and artifact lifecycle"),
            ("configs/pipeline.example.yaml", "Deployed endpoints, GPU/service, routing, parse, structured, and RAG defaults"),
            ("scripts/start_model_services.sh", "Safe model launch/adoption with GPU/model/port checks"),
            ("scripts/run_benchmark.sh", "Bounded validation wrapper and output isolation"),
            ("scripts/monitor_pipeline.py", "Endpoint/GPU/data/run status monitoring"),
            ("docs/operations-and-scale-plan.md", "Transfer, operations, replay, scale constraints, and campaign design"),
        ),
        (3810, 5550),
        font_size=8.35,
    )
    add_title(document, "15.1 Recommended immediate work package", 2)
    add_number(document, "Implement and test the immutable manifest, stable page-window work-unit, and attempt-directory formats.")
    add_number(document, "Refactor page-evidence persistence so independent windows cannot replace one another; merge into a new release database.")
    add_number(document, "Enforce embedding model/dimension and nonempty citation invariants at index and query time.")
    add_number(document, "Run a two-worker vendor pilot with page-level throughput and p95 stage latency instrumentation.")
    add_number(document, "Perform human citation QA on stratified native, table, scan, fallback, blank, rejected, and accepted samples before full vendor execution.")

    add_label(document, "Conclusion")
    add_title(document, "16. Final Assessment", 1)
    add_body(document, "The implemented pipeline is architecturally sound for an evidence-first industrial catalog system. It uses the NVIDIA server where GPU inference adds the most value, keeps deterministic controls around every generative stage, and preserves enough lineage to explain both accepted and rejected outcomes. The hardened replay demonstrates the intended safety behavior: semantic VLM failure did not erase usable native evidence, a genuinely blank page did not waste inference, and an unsupported product did not enter the catalog.")
    add_body(document, "The remaining work is production campaign engineering. The server’s eight A100 GPUs, memory, CPUs, and storage are sufficient to support measured parallelism; however, page coverage, work-unit identity, idempotent merge, and release invariants must be solved before throughput is increased. Once those controls are implemented, the vendor corpus is the correct first full campaign, followed by a separately profiled reference-library campaign and a conflict-checked combined retrieval release.")
    add_callout(document, "Bottom line", "Proceed to the manifest/sharding/merge milestone. Treat the current results as validated architecture and bounded operational proof—not yet as a completed industrial knowledge base.", kind="success")

    add_label(document, "Appendix A")
    add_title(document, "Appendix A. Canonical Product Shape (Conceptual)", 1)
    add_code_block(document, [
        "{",
        '  "manufacturer": {"value": "…", "evidence": ["element_id"]},',
        '  "part_name":    {"value": "…", "evidence": ["element_id"]},',
        '  "part_number":  {"value": "…", "evidence": ["element_id"]},',
        '  "series":       {"value": "…", "evidence": ["element_id"]},',
        '  "specifications": [',
        '    {"name": "…", "value": "…", "unit": "…", "evidence": ["element_id"]}',
        "  ],",
        '  "operating_conditions": [{"name": "…", "value": "…", "evidence": ["element_id"]}],',
        '  "other_details": [{"name": "…", "value": "…", "evidence": ["element_id"]}]',
        "}",
    ])
    add_body(document, "The concrete Pydantic schema is the source of truth. This conceptual shape shows that normalized facts carry evidence references; the validator hydrates those references back to canonical page elements before persistence.")

    add_label(document, "Appendix B")
    add_title(document, "Appendix B. Completion Definitions", 1)
    add_table(
        document,
        ("Term", "Definition"),
        (
            ("Validation batch", "A bounded selection that traverses extraction, structured parsing, validation, storage, embedding, and monitoring. It may include review/failure and is not corpus completion."),
            ("Completed campaign", "Every page in one immutable corpus manifest has a terminal completed/review unit, no unresolved failure, and a verified merged release."),
            ("Completed combined corpus", "Both vendor and library campaigns independently pass their gates, followed by a conflict-checked combined retrieval release."),
        ),
        (2400, 6960),
        font_size=9.0,
    )

    add_label(document, "Appendix C")
    add_title(document, "Appendix C. Authoritative Model References", 1)
    references = [
        ("NVIDIA Nemotron Parse v1.2", "https://huggingface.co/nvidia/NVIDIA-Nemotron-Parse-v1.2", "Document parsing capabilities, required task prompt, vLLM deployment, and A100 support"),
        ("Llama 3.1 Nemotron Nano 8B v1", "https://huggingface.co/nvidia/Llama-3.1-Nemotron-Nano-8B-v1", "Model architecture, long-context behavior, reasoning/chat/RAG/tool-use scope"),
        ("Nemotron 3 Embed 8B BF16", "https://huggingface.co/nvidia/Nemotron-3-Embed-8B-BF16", "Retrieval usage, query/passage prefixes, dimensions, sequence length, and deployment"),
    ]
    for title, url, purpose in references:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title + " — ")
        set_run_font(r, bold=True, color=NAVY)
        add_hyperlink(p, "official NVIDIA model card", url)
        p2 = document.add_paragraph(purpose)
        p2.paragraph_format.left_indent = Inches(0.2)
        p2.paragraph_format.space_after = Pt(7)
        for run in p2.runs:
            set_run_font(run, size=9.5, color=MID_GRAY)
    add_body(document, "Implementation claims and measured results in this report are derived from the local project source, its operational runbook, and the verified server execution artifacts summarized there. Model capability descriptions are limited to the official NVIDIA model cards above.")

    add_label(document, "Appendix D")
    add_title(document, "Appendix D. Glossary", 1)
    add_table(
        document,
        ("Term", "Meaning in this report"),
        (
            ("VLM", "Vision-language model that interprets rendered document pages"),
            ("LLM", "Language model used here for schema-constrained product parsing"),
            ("RAG", "Retrieval-augmented generation; here, a citation-bearing searchable knowledge base"),
            ("NIM-style endpoint", "OpenAI-compatible locally hosted NVIDIA model service used by the clients"),
            ("Source guard", "Deterministic rule layer that rejects manufacturer, cardinality, identifier, or evidence violations"),
            ("Evidence hydration", "Replacement of compact element references with canonical extraction evidence"),
            ("Review", "Completed document with accepted data plus quality warnings; not an execution failure"),
            ("Fail closed", "Stop or reject when evidence is insufficient instead of guessing or silently accepting"),
        ),
        (2500, 6860),
        font_size=8.9,
    )

    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the NVIDIA industrial catalog extraction technical report.")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--asset-dir", type=Path, required=True)
    args = parser.parse_args()
    build_report(args.output.resolve(), args.asset_dir.resolve())
    print(args.output.resolve())


if __name__ == "__main__":
    main()
